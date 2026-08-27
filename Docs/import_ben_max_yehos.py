import psycopg2
import openpyxl
from datetime import datetime

# ==============================================================================
# CONFIGURAÇÃO DE MAPEAMENTO DO IMPORTADOR E EXPORTADOR (BASE DE PRODUÇÃO)
# ==============================================================================
# ID 15 corresponde a 'BENMAX IMPORTACAO E EXPORTACAO LTDA' na base restaurada
IMPORTADOR_ID = 15  

# ID 11 corresponde a 'CHANGSHU LINGKE' na base restaurada (Porto Saída: Shanghai)
EXPORTADOR_ID = 11  

def clean_val(val):
    if val is None:
        return 0.0
    if isinstance(val, str):
        val = val.strip().replace('%', '')
        if val == "" or val.lower() == "none" or val == "-":
            return 0.0
        val = val.replace(',', '.')
        try:
            return float(val)
        except ValueError:
            return val
    return float(val)

def main():
    excel_path = r"c:\Portifólio\pi-web\Docs\EDC - BEN MAX YEHOS 1 - 2 - 3 - 4.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    # 1. Conexão com o banco local
    print("Conectando ao banco de dados...")
    conn = psycopg2.connect("host=localhost port=5432 dbname=pi_db user=pi password=pi123")
    cur = conn.cursor()
    
    try:
        # Iniciando transação
        cur.execute("BEGIN;")
        
        # 2. Garantir que o NCM 84186999 existe no banco
        # Alíquotas conforme planilha: II=12.6%, IPI=9.75%, PIS=2.1%, COFINS=9.65%
        print("Verificando NCM 84186999...")
        cur.execute('SELECT "Id" FROM edc.ncms WHERE "Codigo" = \'84186999\'')
        ncm_row = cur.fetchone()
        if ncm_row:
            ncm_id = ncm_row[0]
            print(f"NCM 84186999 já existe (ID: {ncm_id})")
        else:
            cur.execute("""
                INSERT INTO edc.ncms ("Codigo", "Descricao", "AliquotaII", "AliquotaIPI", "AliquotaPis", "AliquotaCofins", "AliquotaIcmsPadrao", "FlAtivo")
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s) RETURNING "Id"
            """, ('84186999', 'COOLERS / REFRIGERADORES', 0.1260, 0.0975, 0.0210, 0.0965, 0.0260, True))
            ncm_id = cur.fetchone()[0]
            print(f"NCM 84186999 inserido com sucesso (ID: {ncm_id})")

        # 3. Resolver/Garantir Importador (Cliente)
        global IMPORTADOR_ID
        if IMPORTADOR_ID is None:
            cur.execute('SELECT "Id" FROM edc.importadores WHERE "RazaoSocial" = \'BENMAX IMPORTACAO E EXPORTACAO LTDA\'')
            imp_row = cur.fetchone()
            if imp_row:
                IMPORTADOR_ID = imp_row[0]
                print(f"Importador 'BENMAX IMPORTACAO E EXPORTACAO LTDA' encontrado (ID: {IMPORTADOR_ID})")
            else:
                cur.execute("""
                    INSERT INTO edc.importadores ("RazaoSocial", "Cnpj", "UF", "RegimeTributario", "AliquotaIcmsPadrao", "FlAtivo")
                    VALUES (%s, %s, %s, %s, %s, %s) RETURNING "Id"
                """, ('BENMAX IMPORTACAO E EXPORTACAO LTDA', '01.904.966/0003-88', 'SC', 'Lucro Real', 0.0400, True))
                IMPORTADOR_ID = cur.fetchone()[0]
                print(f"Importador cadastrado automaticamente (ID: {IMPORTADOR_ID})")
        else:
            print(f"Utilizando Importador ID configurado: {IMPORTADOR_ID}")
            
        # 4. Resolver/Garantir Exportador
        global EXPORTADOR_ID
        if EXPORTADOR_ID is None:
            cur.execute('SELECT "Id" FROM edc.exportadores WHERE "Nome" = \'CHANGSHU LINGKE\'')
            exp_row = cur.fetchone()
            if exp_row:
                EXPORTADOR_ID = exp_row[0]
                print(f"Exportador 'CHANGSHU LINGKE' encontrado (ID: {EXPORTADOR_ID})")
            else:
                cur.execute("""
                    INSERT INTO edc.exportadores ("Nome", "Pais", "FlAtivo", "Incoterm")
                    VALUES (%s, %s, %s, %s) RETURNING "Id"
                """, ('CHANGSHU LINGKE', 'SHANGHAI', True, 'FOB'))
                EXPORTADOR_ID = cur.fetchone()[0]
                print(f"Exportador cadastrado automaticamente (ID: {EXPORTADOR_ID})")
        else:
            print(f"Utilizando Exportador ID configurado: {EXPORTADOR_ID}")

        # 5. Ler e Inserir Produtos/Modelos da sheet 'LISTA DE COMPRAS'
        print("Lendo itens da planilha 'LISTA DE COMPRAS'...")
        sheet_compras = wb['LISTA DE COMPRAS']
        
        # Mapeamento de itens
        # Rows 6 a 24 contêm os produtos
        parsed_items = []
        for r in range(6, 25):
            ref = sheet_compras.cell(row=r, column=3).value
            if not ref or str(ref).strip() == "" or ref == "TOTAL":
                continue
            
            ref = str(ref).strip()
            qty = clean_val(sheet_compras.cell(row=r, column=10).value)
            fob_unit = clean_val(sheet_compras.cell(row=r, column=11).value)
            
            # Alguns itens têm preço unitário subfaturado
            fob_sub = sheet_compras.cell(row=r, column=12).value
            fob_sub_val = clean_val(fob_sub) if fob_sub is not None else None
            
            parsed_items.append({
                "ref": ref,
                "qty": qty,
                "fob_unit": fob_unit,
                "fob_sub": fob_sub_val
            })
            
        print(f"Encontrados {len(parsed_items)} itens para inserir/mapear.")
        
        # Cadastrar produtos e modelos no banco se não existirem
        item_mappings = []
        for item in parsed_items:
            # Produto
            cur.execute('SELECT "Id" FROM edc.produtos WHERE "Referencia" = %s', (item["ref"],))
            prod_row = cur.fetchone()
            if prod_row:
                prod_id = prod_row[0]
            else:
                cur.execute("""
                    INSERT INTO edc.produtos ("Referencia", "Descricao", "IdNcm", "PesoLiquido", "PesoBruto", "CubagemM3", "UnidadeMedida", "PrecoFobBase", "FlAtivo")
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s) RETURNING "Id"
                """, (item["ref"], item["ref"], ncm_id, 0.0, 0.0, 0.0, 'UN', item["fob_unit"], True))
                prod_id = cur.fetchone()[0]
                print(f"  Produto cadastrado: '{item['ref']}' (ID: {prod_id})")
                
            # Modelo
            cur.execute('SELECT "Id" FROM edc.modelos WHERE "Codigo" = %s AND "IdProduto" = %s', (item["ref"], prod_id))
            mod_row = cur.fetchone()
            if mod_row:
                mod_id = mod_row[0]
            else:
                cur.execute("""
                    INSERT INTO edc.modelos ("Codigo", "Nome", "Descricao", "IdProduto", "FlAtivo")
                    VALUES (%s, %s, %s, %s, %s) RETURNING "Id"
                """, (item["ref"], item["ref"], f"Modelo Padrão - {item['ref']}", prod_id, True))
                mod_id = cur.fetchone()[0]
                print(f"  Modelo cadastrado: '{item['ref']}' (ID: {mod_id})")
                
            item_mappings.append({
                "prod_id": prod_id,
                "mod_id": mod_id,
                "qty": item["qty"],
                "fob_unit": item["fob_unit"],
                "fob_sub": item["fob_sub"]
            })

        # 6. Criar Simulação de Custos (edc.simulacoes)
        print("Inserindo cabeçalho da Simulação...")
        
        # Porto de Saída: SHANGHAI (ID 2), Entrada: ITAPOA (ID 6)
        porto_origem_id = 2
        porto_destino_id = 6
        
        # Cotação e Frete Internacional
        cotacao_dolar = 5.3797  # conforme célula D12
        frete_internacional = 10660.00  # conforme célula F19 (10659.9996...)
        seguro_internacional = 0.00
        
        numero_ref = "EDC - BEN MAX YEHOS 1 - 2 - 3 - 4"
        data_estudo = datetime(2025, 11, 26)  # Conforme célula G9 do Resumo
        
        cur.execute("""
            INSERT INTO edc.simulacoes (
                "NumeroReferencia", "DataEstudo", "IdImportador", "IdExportador", 
                "IdPortoOrigem", "IdPortoDestino", "CotacaoDolar", "SpreadCambio", 
                "TipoFrete", "ValorFreteInternacional", "ValorSeguroInternacional", "Status", 
                "ComissaoPercentual", "FlExibirComissao", "FlSimularSubfaturamento", 
                "MetodoCalculoFederais", "MetodoCalculoIcms", "PercentualSubfaturamento", "ModalidadeFrete"
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s) RETURNING "Id"
        """, (
            numero_ref, data_estudo, IMPORTADOR_ID, EXPORTADOR_ID,
            porto_origem_id, porto_destino_id, cotacao_dolar, 0.00,
            'FOB', frete_internacional, seguro_internacional, 'Rascunho',
            0.0000, False, False, 'SimplificadoExcel', 'SimplificadoExcel', 100.00, '4x40HC'
        ))
        sim_id = cur.fetchone()[0]
        print(f"Simulação criada com ID: {sim_id}")

        # 7. Inserir itens da Simulação (edc.simulacao_itens)
        print("Inserindo itens da Simulação...")
        for im in item_mappings:
            cur.execute("""
                INSERT INTO edc.simulacao_itens (
                    "IdSimulacao", "IdProduto", "Quantidade", "ValorFobUnitario", 
                    "PesoLiquidoTotal", "PesoBrutoTotal", "CubagemTotal", "IdModelo", "ValorFobSubfaturado"
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
            """, (
                sim_id, im["prod_id"], im["qty"], im["fob_unit"],
                0.0, 0.0, 0.0, im["mod_id"], im["fob_sub"]
            ))

        # 8. Ler e Inserir despesas da Simulação (edc.simulacao_despesas)
        print("Inserindo despesas da Simulação...")
        despesas_list = [
            ("TAXA SISCOMEX", 224.00),
            ("ANUENCIA DE LI", 107.06),
            ("DIFERENÇA DO VALOR DO FRETE", 11657.97),
            ("LICENÇA DE IMPORTAÇÃO", 130.00),
            ("AFRMM", 8.00),  # Será computado como 8% do frete
            ("ARMAZENAGEM", 24000.00),
            ("FRETE RODOVIÁRIO", 12200.00),
            ("DESEMBARAÇO ADUANEIRO", 1390.00)
        ]
        
        for idx, (nome, valor) in enumerate(despesas_list):
            cur.execute("""
                INSERT INTO edc.simulacao_despesas ("IdSimulacao", "NomeDespesa", "Valor", "Moeda", "MetodoRateio", "Ordem")
                VALUES (%s, %s, %s, %s, %s, %s)
            """, (sim_id, nome, valor, 'BRL', 'Quantidade', idx))
            print(f"  Despesa inserida: {nome} (Valor: {valor} BRL)")

        # Confirmando a transação
        cur.execute("COMMIT;")
        print("Transação finalizada com sucesso! Todos os dados da EDC foram inseridos no sistema.")
        
    except Exception as e:
        cur.execute("ROLLBACK;")
        print(f"ERRO durante a inserção. Transação abortada: {e}")
        raise e
    finally:
        cur.close()
        conn.close()

if __name__ == "__main__":
    main()
