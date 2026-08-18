import openpyxl
import re
import psycopg2
import unicodedata

def normalize(text):
    if not text:
        return ""
    nfkd_form = unicodedata.normalize('NFKD', str(text))
    text = "".join([c for c in nfkd_form if not unicodedata.combining(c)])
    text = text.upper().strip()
    text = re.sub(r'\s+', ' ', text)
    text = text.replace('´', "'").replace('’', "'").replace('Ç', 'C').replace('Ã', 'A').replace('Õ', 'O')
    text = text.replace('É', 'E').replace('È', 'E').replace('Á', 'A').replace('À', 'A').replace('Ó', 'O')
    text = text.replace('Í', 'I').replace('Ú', 'U').replace('Â', 'A').replace('Ê', 'E').replace('Ô', 'O')
    return text

def clean_val(val):
    if val is None:
        return 0.0
    val_str = str(val).strip().lower()
    if val_str in ("", "none", "-", "0", "0.0"):
        return 0.0
    matches = re.findall(r'(\d+[\.,]\d+|\d+)', val_str)
    if matches:
        return float(matches[0].replace(',', '.'))
    return 0.0

def extract_width(comp_str):
    if comp_str is None:
        return None
    comp_str = str(comp_str).strip()
    matches = re.findall(r'(\d+[\.,]\d+|\d+)', comp_str)
    if matches:
        return float(matches[-1].replace(',', '.'))
    return None

def main():
    conn = psycopg2.connect("host=localhost port=5432 dbname=pi_db user=pi password=pi123")
    cur = conn.cursor()
    
    cur.execute("SELECT id, nome FROM pi.marca;")
    brand_map = {normalize(r[1]): r[0] for r in cur.fetchall()}
    
    cur.execute("""
        SELECT m.id, ma.nome, m.descricao, m.largura, m.profundidade, m.altura, m.m3
        FROM pi.modulo m
        JOIN pi.marca ma ON m.id_marca = ma.id
        WHERE m.id_fornecedor = 4;
    """)
    db_list = []
    for r in cur.fetchall():
        db_list.append({
            "id": r[0],
            "brand_name": normalize(r[1]),
            "descricao_norm": normalize(r[2]),
            "largura": float(r[3]),
            "profundidade": float(r[4])
        })
        
    wb = openpyxl.load_workbook("Docs/TABELA LIVINTUS_EXPORTAÇÃO_MOVELSUL_2026.xlsx", data_only=True)
    sheet = wb['LIVINTUS COURO']
    liv_modules = []
    current_brand = None
    for r in range(3, sheet.max_row + 1):
        brand_val = sheet.cell(row=r, column=2).value
        if brand_val and str(brand_val).strip() and str(brand_val).strip() != "-":
            current_brand = str(brand_val).strip()
            
        comp_val = sheet.cell(row=r, column=3).value
        if comp_val is not None:
            alt_val = sheet.cell(row=r, column=4).value
            prof_val = sheet.cell(row=r, column=5).value
            desc_val = sheet.cell(row=r, column=6).value
            
            width = extract_width(comp_val)
            depth = clean_val(prof_val)
            height = clean_val(alt_val)
            desc = str(desc_val).strip() if desc_val else ""
            
            liv_modules.append({
                "brand": current_brand,
                "desc": desc,
                "width": width,
                "depth": depth,
                "height": height,
                "row": r
            })
            
    print("=== Unmatched items for Livintus ===")
    unmatched_count = 0
    for p in liv_modules:
        p_brand_norm = normalize(p["brand"])
        p_desc_norm = normalize(p["desc"])
        p_w = p["width"]
        p_d = p["depth"]
        
        found = False
        for db in db_list:
            if db["brand_name"] == p_brand_norm:
                # check exact
                if db["descricao_norm"] == p_desc_norm and p_w is not None and abs(db["largura"] - p_w) < 0.02 and p_d is not None and abs(db["profundidade"] - p_d) < 0.02:
                    found = True
                    break
                # check dimension match only (which we match manually or by script logic)
                if p_w is not None and abs(db["largura"] - p_w) < 0.02 and p_d is not None and abs(db["profundidade"] - p_d) < 0.02:
                    found = True
                    break
                # check manual matches for recliners
                if p["brand"].upper() in ("POLTRONA ALANA ELÉTRICA", "CONDOR") and db["descricao_norm"] == p_desc_norm and abs(db["largura"] - p_w) < 0.02:
                    found = True
                    break
        if not found:
            unmatched_count += 1
            print(f"Row {p['row']:3d} | Brand: {p['brand']:<25} | Desc: {p['desc']:<30} | W={p['width']} D={p['depth']} H={p['height']}")
    print(f"Total unmatched: {unmatched_count}")
    
    cur.close()
    conn.close()

if __name__ == "__main__":
    main()
