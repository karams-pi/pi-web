import openpyxl
import re
import psycopg2

def clean_val(val):
    if val is None:
        return 0.0
    val_str = str(val).strip().lower()
    if val_str in ("", "none", "-", "0", "0.0"):
        return 0.0
    matches = re.findall(r'(\d+[\.,]\d+|\d+)', val_str)
    if matches:
        return float(matches[0].replace(',', '.'))
    return 0.0

def extract_width(comp_str):
    if comp_str is None:
        return None
    comp_str = str(comp_str).strip()
    matches = re.findall(r'(\d+[\.,]\d+|\d+)', comp_str)
    if matches:
        return float(matches[-1].replace(',', '.'))
    return None

def normalize(text):
    if not text:
        return ""
    import unicodedata
    nfkd_form = unicodedata.normalize('NFKD', str(text))
    text = "".join([c for c in nfkd_form if not unicodedata.combining(c)])
    text = text.upper().strip()
    text = re.sub(r'\s+', ' ', text)
    text = text.replace('´', "'").replace('’', "'").replace('Ç', 'C').replace('Ã', 'A').replace('Õ', 'O')
    return text

conn = psycopg2.connect("host=localhost port=5432 dbname=pi_db user=pi password=pi123")
cur = conn.cursor()

# Load DB modules for Livintus
cur.execute("""
    SELECT m.id, ma.nome, m.descricao, m.largura, m.profundidade, m.altura, m.m3
    FROM pi.modulo m
    JOIN pi.marca ma ON m.id_marca = ma.id
    WHERE m.id_fornecedor = 4;
""")
db_list = []
for r in cur.fetchall():
    db_list.append({
        "id": r[0],
        "brand_name": normalize(r[1]),
        "descricao_norm": normalize(r[2]),
        "largura": float(r[3]),
        "profundidade": float(r[4])
    })

wb = openpyxl.load_workbook("Docs/TABELA LIVINTUS_EXPORTAÇÃO_MOVELSUL_2026.xlsx", data_only=True)
sheet = wb['LIVINTUS COURO']
current_brand = None

for r in range(3, sheet.max_row + 1):
    brand_val = sheet.cell(row=r, column=2).value
    if brand_val and str(brand_val).strip() and str(brand_val).strip() != "-":
        current_brand = str(brand_val).strip()
        
    comp_val = sheet.cell(row=r, column=3).value
    if comp_val is not None:
        width = extract_width(comp_val)
        prof_val = sheet.cell(row=r, column=5).value
        depth = clean_val(prof_val)
        desc_val = sheet.cell(row=r, column=6).value
        desc = str(desc_val).strip() if desc_val else ""
        
        if r == 66:
            p_brand_norm = normalize(current_brand)
            p_desc_norm = normalize(desc)
            print(f"DEBUG r=66 | brand={current_brand} | desc={desc} | width={width} | depth={depth}")
            print(f"Normalized | brand={p_brand_norm} | desc={p_desc_norm}")
            for db in db_list:
                if db["brand_name"] == p_brand_norm:
                    print(f"Compare with DB ID {db['id']}: desc_norm={db['descricao_norm']} | L={db['largura']} | P={db['profundidade']}")
                    # exact match check
                    exact_ok = db["descricao_norm"] == p_desc_norm and abs(db["largura"] - width) < 0.02 and abs(db["profundidade"] - depth) < 0.02
                    print(f"  Exact check: {exact_ok}")
                    # desc-only check
                    desc_ok = db["descricao_norm"] == p_desc_norm
                    print(f"  Desc check: {desc_ok}")
cur.close()
conn.close()
