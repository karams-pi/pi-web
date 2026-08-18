import openpyxl
import re

def extract_width(comp_str):
    if comp_str is None:
        return None
    comp_str = str(comp_str).strip()
    matches = re.findall(r'(\d+[\.,]\d+|\d+)', comp_str)
    if matches:
        return float(matches[-1].replace(',', '.'))
    return None

wb = openpyxl.load_workbook("Docs/TABELA LIVINTUS_EXPORTAÇÃO_MOVELSUL_2026.xlsx", data_only=True)
sheet = wb['LIVINTUS COURO']
for r in range(60, 75):
    val = sheet.cell(row=r, column=3).value
    brand = sheet.cell(row=r, column=2).value
    prof = sheet.cell(row=r, column=5).value
    desc = sheet.cell(row=r, column=6).value
    if val:
        w = extract_width(val)
        print(f"Row {r} | Brand: {brand} | C: {repr(val)} -> W={w} | P: {repr(prof)} | F: {repr(desc)}")
