import openpyxl

wb = openpyxl.load_workbook("Docs/TABELA FERGUILE_EXPORTAÇÃO_MOVELSUL_2026.xlsx", data_only=True)
sheet = wb['EXW']
for r in range(155, 175):
    brand = sheet.cell(row=r, column=2).value
    comp = sheet.cell(row=r, column=3).value
    prof = sheet.cell(row=r, column=4).value
    alt = sheet.cell(row=r, column=5).value
    desc = sheet.cell(row=r, column=12).value
    if comp is not None or brand is not None:
        print(f"Row {r} | Brand: {brand} | C: {repr(comp)} | P: {prof} | A: {alt} | D: {repr(desc)}")
