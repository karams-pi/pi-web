import openpyxl
import re
import psycopg2
import unicodedata

def normalize(text):
    if not text:
        return ""
    nfkd_form = unicodedata.normalize('NFKD', str(text))
    text = "".join([c for c in nfkd_form if not unicodedata.combining(c)])
    text = text.upper().strip()
    text = re.sub(r'\s+', ' ', text)
    text = text.replace('´', "'").replace('’', "'").replace('Ç', 'C').replace('Ã', 'A').replace('Õ', 'O')
    text = text.replace('É', 'E').replace('È', 'E').replace('Á', 'A').replace('À', 'A').replace('Ó', 'O')
    text = text.replace('Í', 'I').replace('Ú', 'U').replace('Â', 'A').replace('Ê', 'E').replace('Ô', 'O')
    return text

def clean_val(val):
    if val is None:
        return 0.0
    val_str = str(val).strip().lower()
    if val_str in ("", "none", "-", "0", "0.0"):
        return 0.0
    matches = re.findall(r'(\d+[\.,]\d+|\d+)', val_str)
    if matches:
        return float(matches[0].replace(',', '.'))
    return 0.0

def extract_width(comp_str):
    if comp_str is None:
        return None
    comp_str = str(comp_str).strip()
    matches = re.findall(r'(\d+[\.,]\d+|\d+)', comp_str)
    if matches:
        return float(matches[-1].replace(',', '.'))
    return None

def parse_depth(prof_str):
    if prof_str is None:
        return None
    prof_str = str(prof_str).strip().replace(',', '.')
    if "f:" in prof_str.lower():
        m = re.search(r'f:\s*([\d.]+)', prof_str.lower())
        if m:
            return float(m.group(1))
    return clean_val(prof_str)

def main():
    conn = psycopg2.connect("host=localhost port=5432 dbname=pi_db user=pi password=pi123")
    cur = conn.cursor()
    
    # Cache brand name mapping
    cur.execute("SELECT id, nome FROM pi.marca;")
    brand_map = {normalize(r[1]): r[0] for r in cur.fetchall()}
    
    # Load Ferguile DB modules (forn 3)
    cur.execute("""
        SELECT m.id, ma.nome, m.descricao, m.largura, m.profundidade, m.altura, m.m3, m.id_marca
        FROM pi.modulo m
        JOIN pi.marca ma ON m.id_marca = ma.id
        WHERE m.id_fornecedor = 3;
    """)
    db_ferguile = []
    for r in cur.fetchall():
        db_ferguile.append({
            "id": r[0],
            "brand_name": normalize(r[1]),
            "descricao_raw": r[2],
            "descricao_norm": normalize(r[2]),
            "largura": float(r[3]),
            "profundidade": float(r[4]),
            "altura": float(r[5]),
            "m3": float(r[6]),
            "id_marca": r[7]
        })
        
    # Load Livintus DB modules (forn 4)
    cur.execute("""
        SELECT m.id, ma.nome, m.descricao, m.largura, m.profundidade, m.altura, m.m3, m.id_marca
        FROM pi.modulo m
        JOIN pi.marca ma ON m.id_marca = ma.id
        WHERE m.id_fornecedor = 4;
    """)
    db_livintus = []
    for r in cur.fetchall():
        db_livintus.append({
            "id": r[0],
            "brand_name": normalize(r[1]),
            "descricao_raw": r[2],
            "descricao_norm": normalize(r[2]),
            "largura": float(r[3]),
            "profundidade": float(r[4]),
            "altura": float(r[5]),
            "m3": float(r[6]),
            "id_marca": r[7]
        })

    def find_match(p, db_list):
        p_brand_norm = normalize(p["brand"])
        p_desc_norm = normalize(p["desc"])
        p_w = p["width"]
        p_d = p["depth"]
        p_h = p["height"]
        
        # 1. Exact match (brand + desc + width + depth) within 2cm
        for db in db_list:
            if db["brand_name"] == p_brand_norm:
                if db["descricao_norm"] == p_desc_norm:
                    if p_w is not None and abs(db["largura"] - p_w) < 0.02 and p_d is not None and abs(db["profundidade"] - p_d) < 0.02:
                        return db, "Exact Match"
                        
        # 2. Dimension match only (brand + width + depth) within 2cm
        for db in db_list:
            if db["brand_name"] == p_brand_norm:
                if p_w is not None and abs(db["largura"] - p_w) < 0.02 and p_d is not None and abs(db["profundidade"] - p_d) < 0.02:
                    return db, "Dimension Match (diff desc)"
                    
        # 3. Swapped dimensions match (brand + width + depth/height swapped) within 2cm
        for db in db_list:
            if db["brand_name"] == p_brand_norm:
                if p_w is not None and abs(db["largura"] - p_w) < 0.02:
                    # check if db depth matches sheet height and db height matches sheet depth
                    if p_h is not None and abs(db["profundidade"] - p_h) < 0.02 and p_d is not None and abs(db["altura"] - p_d) < 0.02:
                        return db, "Swapped Dimensions Match"
                        
        # 4. Recliner description match (brand + description + width) within 2cm (handles different depth when opened)
        for db in db_list:
            if db["brand_name"] == p_brand_norm:
                if db["descricao_norm"] == p_desc_norm:
                    if p_w is not None and abs(db["largura"] - p_w) < 0.02:
                        return db, "Recliner Match (description + width)"
                        
        # 5. Insieme width-based match (brand = INSIEME + width) within 2cm
        if p_brand_norm == "INSIEME":
            for db in db_list:
                if db["brand_name"] == "INSIEME":
                    if p_w is not None and abs(db["largura"] - p_w) < 0.02:
                        return db, "Insieme Width Match"
                        
        return None, "No Match"

    # Process Ferguile
    wb_fer = openpyxl.load_workbook("Docs/TABELA FERGUILE_EXPORTAÇÃO_MOVELSUL_2026.xlsx", data_only=True)
    sheet_fer = wb_fer['EXW']
    current_brand = None
    fer_parsed = []
    for r in range(3, sheet_fer.max_row + 1):
        brand_val = sheet_fer.cell(row=r, column=2).value
        if brand_val and str(brand_val).strip() and str(brand_val).strip() != "-":
            current_brand = str(brand_val).strip()
        comp_val = sheet_fer.cell(row=r, column=3).value
        if comp_val is not None:
            prof_val = sheet_fer.cell(row=r, column=4).value
            alt_val = sheet_fer.cell(row=r, column=5).value
            desc_val = sheet_fer.cell(row=r, column=12).value
            
            p = {
                "brand": current_brand,
                "desc": str(desc_val).strip() if desc_val else "",
                "width": extract_width(comp_val),
                "depth": parse_depth(prof_val),
                "height": clean_val(alt_val),
                "row": r
            }
            db_match, match_type = find_match(p, db_ferguile)
            print(f"Ferguile Row {r:3d} | Brand: {p['brand']:<15} | Desc: {p['desc']:<40} | W={p['width']} D={p['depth']} -> Match: {match_type} (DB ID: {db_match['id'] if db_match else 'None'})")

    # Process Livintus
    wb_liv = openpyxl.load_workbook("Docs/TABELA LIVINTUS_EXPORTAÇÃO_MOVELSUL_2026.xlsx", data_only=True)
    sheet_liv = wb_liv['LIVINTUS COURO']
    current_brand = None
    for r in range(3, sheet_liv.max_row + 1):
        brand_val = sheet_liv.cell(row=r, column=2).value
        if brand_val and str(brand_val).strip():
            # handle ditto mark '-'
            brand_str = str(brand_val).strip()
            if brand_str != "-":
                current_brand = brand_str
        comp_val = sheet_liv.cell(row=r, column=3).value
        if comp_val is not None:
            alt_val = sheet_liv.cell(row=r, column=4).value
            prof_val = sheet_liv.cell(row=r, column=5).value
            desc_val = sheet_liv.cell(row=r, column=6).value
            
            p = {
                "brand": current_brand,
                "desc": str(desc_val).strip() if desc_val else "",
                "width": extract_width(comp_val),
                "depth": clean_val(prof_val),
                "height": clean_val(alt_val),
                "row": r
            }
            db_match, match_type = find_match(p, db_livintus)
            print(f"Livintus Row {r:3d} | Brand: {p['brand']:<15} | Desc: {p['desc']:<40} | W={p['width']} D={p['depth']} -> Match: {match_type} (DB ID: {db_match['id'] if db_match else 'None'})")

    cur.close()
    conn.close()

if __name__ == "__main__":
    main()
