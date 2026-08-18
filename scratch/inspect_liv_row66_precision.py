import openpyxl
import re

def clean_val(val):
    if val is None:
        return 0.0
    val_str = str(val).strip().lower()
    if val_str in ("", "none", "-", "0", "0.0"):
        return 0.0
    matches = re.findall(r'(\d+[\.,]\d+|\d+)', val_str)
    if matches:
        return float(matches[0].replace(',', '.'))
    return 0.0

def extract_width(comp_str):
    if comp_str is None:
        return None
    comp_str = str(comp_str).strip()
    matches = re.findall(r'(\d+[\.,]\d+|\d+)', comp_str)
    if matches:
        return float(matches[-1].replace(',', '.'))
    return None

wb = openpyxl.load_workbook("Docs/TABELA LIVINTUS_EXPORTAÇÃO_MOVELSUL_2026.xlsx", data_only=True)
sheet = wb['LIVINTUS COURO']

liv_modules = []
current_brand = None
for r in range(3, sheet.max_row + 1):
    brand_val = sheet.cell(row=r, column=2).value
    if brand_val and str(brand_val).strip() and str(brand_val).strip() != "-":
        current_brand = str(brand_val).strip()
        
    comp_val = sheet.cell(row=r, column=3).value
    if comp_val is not None:
        alt_val = sheet.cell(row=r, column=4).value
        prof_val = sheet.cell(row=r, column=5).value
        desc_val = sheet.cell(row=r, column=6).value
        
        width = extract_width(comp_val)
        depth = clean_val(prof_val)
        height = clean_val(alt_val)
        desc = str(desc_val).strip() if desc_val else ""
        
        m3 = 0.0
        if isinstance(width, float) and isinstance(depth, float) and isinstance(height, float):
            m3 = round(width * depth * height, 4)
            
        liv_modules.append({
            "brand": current_brand,
            "desc": desc,
            "width": width,
            "depth": depth,
            "height": height,
            "m3": m3,
            "row": r
        })

for m in liv_modules:
    if m["row"] == 66:
        print("Parsed row 66 module:", m)
