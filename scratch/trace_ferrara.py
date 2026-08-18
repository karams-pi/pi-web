import openpyxl
import re
import psycopg2

def clean_val(val):
    if val is None:
        return 0.0
    val_str = str(val).strip().lower()
    if val_str in ("", "none", "-", "0", "0.0"):
        return 0.0
    matches = re.findall(r'(\d+[\.,]\d+|\d+)', val_str)
    if matches:
        return float(matches[0].replace(',', '.'))
    return 0.0

def extract_width(comp_str):
    if comp_str is None:
        return None
    comp_str = str(comp_str).strip()
    matches = re.findall(r'(\d+[\.,]\d+|\d+)', comp_str)
    if matches:
        return float(matches[-1].replace(',', '.'))
    return None

def parse_depth(prof_str):
    if prof_str is None:
        return None
    prof_str = str(prof_str).strip().replace(',', '.')
    if "f:" in prof_str.lower():
        m = re.search(r'f:\s*([\d.]+)', prof_str.lower())
        if m:
            return float(m.group(1))
    return clean_val(prof_str)

def normalize(text):
    if not text:
        return ""
    import unicodedata
    nfkd_form = unicodedata.normalize('NFKD', str(text))
    text = "".join([c for c in nfkd_form if not unicodedata.combining(c)])
    text = text.upper().strip()
    text = re.sub(r'\s+', ' ', text)
    text = text.replace('´', "'").replace('’', "'").replace('Ç', 'C').replace('Ã', 'A').replace('Õ', 'O')
    return text

conn = psycopg2.connect("host=localhost port=5432 dbname=pi_db user=pi password=pi123")
cur = conn.cursor()

cur.execute("""
    SELECT m.id, ma.nome, m.descricao, m.largura, m.profundidade, m.altura, m.m3
    FROM pi.modulo m
    JOIN pi.marca ma ON m.id_marca = ma.id
    WHERE m.id_fornecedor = 3 AND UPPER(ma.nome) = 'FERRARA';
""")
db_list = []
for r in cur.fetchall():
    db_list.append({
        "id": r[0],
        "brand_name": normalize(r[1]),
        "descricao_norm": normalize(r[2]),
        "largura": float(r[3]),
        "profundidade": float(r[4])
    })

wb = openpyxl.load_workbook("Docs/TABELA FERGUILE_EXPORTAÇÃO_MOVELSUL_2026.xlsx", data_only=True)
sheet = wb['EXW']
for r in range(150, 175):
    comp_val = sheet.cell(row=r, column=3).value
    if comp_val is not None:
        brand_val = sheet.cell(row=r, column=2).value
        prof_val = sheet.cell(row=r, column=4).value
        desc_val = sheet.cell(row=r, column=12).value
        
        width = extract_width(comp_val)
        depth = parse_depth(prof_val)
        desc = str(desc_val).strip() if desc_val else ""
        
        if "FERRARA" in str(brand_val).upper() or r in (157, 163, 169):
            print(f"Row {r} | C={repr(comp_val)} -> W={width} | P={repr(prof_val)} -> D={depth} | D={repr(desc)}")
            p_brand_norm = "FERRARA"
            p_desc_norm = normalize(desc)
            for db in db_list:
                print(f"  DB ID {db['id']}: desc_norm={db['descricao_norm']} | L={db['largura']} | P={db['profundidade']}")
                exact_ok = db["descricao_norm"] == p_desc_norm and abs(db["largura"] - width) < 0.02 and abs(db["profundidade"] - depth) < 0.02
                print(f"    Exact match check: {exact_ok}")

cur.close()
conn.close()
