import openpyxl

def inspect_livintus():
    file_path = r"Docs/TABELA LIVINTUS_EXPORTAÇÃO_MOVELSUL_2026.xlsx"
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb['LIVINTUS COURO']
    print(f"Sheet LIVINTUS COURO (Rows: {sheet.max_row}, Cols: {sheet.max_column})")
    
    for r in range(1, 100):
        row_vals = []
        has_val = False
        for c in range(1, sheet.max_column + 1):
            val = sheet.cell(row=r, column=c).value
            if val is not None:
                has_val = True
                if isinstance(val, str):
                    val = val.replace('\n', ' ').replace('\r', ' ')
                row_vals.append(f"{openpyxl.utils.get_column_letter(c)}{r}: {val}")
            else:
                row_vals.append(f"{openpyxl.utils.get_column_letter(c)}{r}: None")
        if has_val:
            while len(row_vals) > 0 and row_vals[-1].endswith(": None"):
                row_vals.pop()
            print(f"Row {r:03d}: " + " | ".join(row_vals))

if __name__ == "__main__":
    inspect_livintus()
