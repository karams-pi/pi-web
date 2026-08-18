import openpyxl
import re

def clean_val(val):
    if val is None:
        return 0.0
    val_str = str(val).strip().lower()
    if val_str in ("", "none", "-", "0", "0.0"):
        return 0.0
    matches = re.findall(r'(\d+[\.,]\d+|\d+)', val_str)
    if matches:
        return float(matches[0].replace(',', '.'))
    return 0.0

def extract_width(comp_str):
    if comp_str is None:
        return None
    comp_str = str(comp_str).strip()
    matches = re.findall(r'(\d+[\.,]\d+|\d+)', comp_str)
    if matches:
        return float(matches[-1].replace(',', '.'))
    return None

wb = openpyxl.load_workbook("Docs/TABELA LIVINTUS_EXPORTAÇÃO_MOVELSUL_2026.xlsx", data_only=True)
sheet = wb['LIVINTUS COURO']

for r in range(3, sheet.max_row + 1):
    comp_val = sheet.cell(row=r, column=3).value
    if comp_val is not None:
        width = extract_width(comp_val)
        if r == 66:
            print(f"Row {r} | comp_val={repr(comp_val)} | width={width}")
            prof_val = sheet.cell(row=r, column=5).value
            depth = clean_val(prof_val)
            print(f"prof_val={repr(prof_val)} | depth={depth}")
