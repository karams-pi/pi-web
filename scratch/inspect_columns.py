import openpyxl

def main():
    f1 = "Docs/TABELA FERGUILE_EXPORTAÇÃO_MOVELSUL_2026.xlsx"
    wb1 = openpyxl.load_workbook(f1, data_only=True)
    sheet1 = wb1['EXW']
    print("--- Ferguile EXW Headers ---")
    for c in range(1, sheet1.max_column + 1):
        h = sheet1.cell(row=2, column=c).value
        print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): {h}")
        
    f2 = "Docs/TABELA LIVINTUS_EXPORTAÇÃO_MOVELSUL_2026.xlsx"
    wb2 = openpyxl.load_workbook(f2, data_only=True)
    sheet2 = wb2['LIVINTUS COURO']
    print("\n--- Livintus Couro Headers ---")
    for c in range(1, sheet2.max_column + 1):
        h = sheet2.cell(row=2, column=c).value
        print(f"Col {c} ({openpyxl.utils.get_column_letter(c)}): {h}")

if __name__ == "__main__":
    main()
