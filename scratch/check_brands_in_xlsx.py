import openpyxl

def check_file(file_path, sheet_name):
    print(f"\nChecking: {file_path} [{sheet_name}]")
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb[sheet_name]
    
    current_brand = None
    for r in range(1, sheet.max_row + 1):
        a_val = sheet.cell(row=r, column=1).value
        b_val = sheet.cell(row=r, column=2).value
        
        # Print non-empty rows for Col A and B
        if a_val or b_val:
            print(f"Row {r:03d} | A: {a_val} | B: {b_val}")

if __name__ == "__main__":
    check_file("Docs/TABELA FERGUILE_EXPORTAÇÃO_MOVELSUL_2026.xlsx", "EXW")
    check_file("Docs/TABELA LIVINTUS_EXPORTAÇÃO_MOVELSUL_2026.xlsx", "LIVINTUS COURO")
